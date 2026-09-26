"""Build the final SPORTSGD academic report from docs/REPORT_CONTENT.md.

The script intentionally supports the Markdown subset used by the report source:
headings, paragraphs, lists, block quotes, pipe tables, fenced code, images and
inline bold/italic/code/links. Word updates the generated TOC fields during the
PDF export step documented in README/IMPLEMENTATION_REPORT.
"""

from __future__ import annotations

import argparse
import re
import textwrap
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "REPORT_CONTENT.md"
OUTPUT = ROOT / "docs" / "SPORTSGD_Reporte_Final.docx"

INK = RGBColor(0x1E, 0x1E, 0x1E)
MUTED = RGBColor(0x5E, 0x5A, 0x5A)
BRAND = RGBColor(0x00, 0x7F, 0x6D)
SURFACE = "E6F6F4"
LIGHT_GRAY = "F2F2F2"


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    """Apply the restrained horizontal rules recommended for APA tables."""
    table_properties = table._tbl.tblPr
    existing = table_properties.find(qn("w:tblBorders"))
    if existing is not None:
        table_properties.remove(existing)

    borders = OxmlElement("w:tblBorders")
    for edge, value in (
        ("top", "single"),
        ("left", "nil"),
        ("bottom", "single"),
        ("right", "nil"),
        ("insideH", "nil"),
        ("insideV", "nil"),
    ):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), value)
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    table_properties.append(borders)


def set_cell_bottom_border(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    existing = cell_properties.find(qn("w:tcBorders"))
    if existing is not None:
        cell_properties.remove(existing)
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    cell_properties.append(borders)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = "Actualizar campo") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_node, separate, text_node, end])


def add_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "007F6D")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|https?://[^\s)]+|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
)


def add_inline(paragraph, text: str, *, base_size: float | None = None) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            if base_size:
                run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("http"):
            add_hyperlink(paragraph, token, token)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if base_size:
                run.font.size = Pt(base_size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size or 10)
            run.font.color.rgb = BRAND
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            if base_size:
                run.font.size = Pt(base_size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        if base_size:
            run.font.size = Pt(base_size)


def strip_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    return text.strip()


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for style_name, size, bold, italic, alignment in (
        ("Title", 20, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Subtitle", 14, False, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 14, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, True, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, True, True, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = INK
        style.paragraph_format.alignment = alignment
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    for style_name in ("Figure Caption", "Table Caption"):
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = INK
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    if "Figure Note" not in styles:
        styles.add_style("Figure Note", WD_STYLE_TYPE.PARAGRAPH)
    note = styles["Figure Note"]
    note.font.name = "Times New Roman"
    note.font.size = Pt(10)
    note.font.color.rgb = MUTED
    note.paragraph_format.first_line_indent = Inches(0)
    note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    note.paragraph_format.space_after = Pt(0)

    if "Code Block" not in styles:
        styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.keep_together = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.5)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_field(paragraph, " PAGE ", "1")
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    core = document.core_properties
    core.title = "SPORTSGD - Reporte técnico final"
    core.subject = "Desarrollo de Aplicaciones Móviles"
    core.author = "Emiliano Iturralde Velazquez; Antonio de Jesus Juarez Padilla"
    core.keywords = "SPORTSGD, Android, Kotlin, Room, Figma, UX"
    core.comments = "Reporte académico final generado a partir de evidencia verificada."


def add_cover(document: Document) -> None:
    for _ in range(2):
        document.add_paragraph()
    university = document.add_paragraph()
    university.alignment = WD_ALIGN_PARAGRAPH.CENTER
    university.paragraph_format.first_line_indent = Inches(0)
    run = university.add_run("Universidad Tecmilenio")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    document.add_paragraph()
    report = document.add_paragraph()
    report.style = document.styles["Subtitle"]
    report.add_run("Proyecto")

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.add_run("SPORTSGD")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Inches(0)
    subtitle.add_run("Reporte técnico final").italic = True

    document.add_paragraph()
    data = document.add_table(rows=0, cols=1)
    data.alignment = WD_TABLE_ALIGNMENT.CENTER
    data.autofit = True
    entries = [
        ("Materia", "Desarrollo de Aplicaciones Móviles"),
        ("Profesor", "Jose Luis Suchil Miranda"),
        (
            "Autores",
            "Emiliano Iturralde Velazquez\nAntonio de Jesus Juarez Padilla",
        ),
        ("Fecha de entrega", "25 de septiembre de 2026"),
    ]
    for label, value in entries:
        cell = data.add_row().cells[0]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        label_run = paragraph.add_run(f"{label}:\n")
        label_run.bold = True
        label_run.font.name = "Times New Roman"
        label_run.font.size = Pt(12)
        value_run = paragraph.add_run(value)
        value_run.font.name = "Times New Roman"
        value_run.font.size = Pt(12)
    document.add_page_break()


def add_indexes(document: Document) -> None:
    heading = document.add_heading("Índice general", level=1)
    heading.paragraph_format.keep_with_next = True
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Inches(0)
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "Actualice el índice en Word")

    document.add_paragraph()
    figure_heading = document.add_heading("Índice de figuras", level=2)
    figure_heading.paragraph_format.keep_with_next = True
    figure_toc = document.add_paragraph()
    figure_toc.paragraph_format.first_line_indent = Inches(0)
    add_field(
        figure_toc,
        ' TOC \\h \\z \\t "Figure Caption,1" ',
        "Actualice el índice de figuras en Word",
    )

    document.add_paragraph()
    table_heading = document.add_heading("Índice de tablas", level=2)
    table_heading.paragraph_format.keep_with_next = True
    table_toc = document.add_paragraph()
    table_toc.paragraph_format.first_line_indent = Inches(0)
    add_field(
        table_toc,
        ' TOC \\h \\z \\t "Table Caption,1" ',
        "Actualice el índice de tablas en Word",
    )
    document.add_page_break()


def resolve_image(markdown_path: str) -> Path:
    return (SOURCE.parent / markdown_path).resolve()


def add_image(document: Document, path: Path, alt: str) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Missing report image: {path}")
    with Image.open(path) as image:
        width_px, height_px = image.size
    max_width = 6.15
    max_height = 6.15
    aspect = width_px / max(height_px, 1)
    width = min(max_width, max_height * aspect)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    picture = run._r.xpath(".//wp:docPr")
    if picture:
        picture[0].set("descr", alt)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for index in range(columns):
        value = rows[0][index] if index < len(rows[0]) else ""
        cell = header.cells[index]
        set_cell_margins(cell)
        set_cell_bottom_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(strip_inline(value))
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)
    prevent_row_split(header)

    for source_row in rows[1:]:
        row = table.add_row()
        prevent_row_split(row)
        for index in range(columns):
            value = source_row[index] if index < len(source_row) else ""
            cell = row.cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, base_size=8.5)

    # Keep compact tables on one page when space permits. If a table is taller
    # than a page, Word still breaks it safely between non-splittable rows and
    # repeats the header instead of leaving a caption plus an orphaned header.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    trailing = document.add_paragraph()
    trailing.paragraph_format.first_line_indent = Inches(0)
    trailing.paragraph_format.space_after = Pt(0)


def parse_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_caption(document: Document, label: str, title: str, style: str) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run().add_break()
    title_run = paragraph.add_run(title)
    title_run.italic = True


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shading)
    for line_index, line in enumerate(lines):
        wrapped = textwrap.wrap(
            line,
            width=105,
            replace_whitespace=False,
            drop_whitespace=False,
            subsequent_indent="    ",
        ) or [""]
        for part_index, part in enumerate(wrapped):
            run = paragraph.add_run(part)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            if line_index != len(lines) - 1 or part_index != len(wrapped) - 1:
                run.add_break()


def parse_markdown(document: Document, lines: list[str]) -> None:
    index = 0
    in_code = False
    code_lines: list[str] = []
    first_h1 = True
    in_references = False
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(piece.strip() for piece in paragraph_buffer).strip()
        if text:
            paragraph = document.add_paragraph()
            if in_references:
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
                paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text)
        paragraph_buffer = []

    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("<!--"):
            flush_paragraph()
            # Major sections already receive deterministic page breaks. Ignoring
            # source layout comments prevents consecutive breaks and blank pages.
            index += 1
            continue
        if stripped == "---":
            flush_paragraph()
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = strip_inline(heading.group(2))
            if level == 1 and not first_h1:
                document.add_page_break()
            if level == 1:
                first_h1 = False
                in_references = title == "Referencias"
            paragraph = document.add_heading(title, level=level)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            index += 1
            continue

        caption = re.fullmatch(r"\*\*((?:Figura|Tabla)\s+\d+)\*\*\s*", stripped)
        if caption:
            flush_paragraph()
            label = caption.group(1)
            title = ""
            if index + 1 < len(lines):
                title_line = lines[index + 1].strip()
                if title_line.startswith("*") and title_line.endswith("*"):
                    title = strip_inline(title_line)
                    index += 1
            style = "Figure Caption" if label.startswith("Figura") else "Table Caption"
            add_caption(document, label, title, style)
            index += 1
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            flush_paragraph()
            add_image(document, resolve_image(image.group(2)), image.group(1))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            flush_paragraph()
            table_rows = [parse_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(document, table_rows)
            continue

        list_item = re.match(r"^(\s*)([-*]|\d+\.)\s+(.+)$", raw)
        if list_item:
            flush_paragraph()
            indentation = len(list_item.group(1)) // 3
            ordered = list_item.group(2).endswith(".")
            if ordered:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25 + 0.25 * indentation)
                paragraph.paragraph_format.first_line_indent = Inches(-0.2)
                add_inline(paragraph, f"{list_item.group(2)} {list_item.group(3)}")
            else:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.25 * indentation)
                paragraph.paragraph_format.first_line_indent = Inches(0)
                add_inline(paragraph, list_item.group(3))
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            paragraph.paragraph_format.space_after = Pt(0)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_parts: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_parts.append(lines[index].strip()[1:].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            paragraph.paragraph_format.space_after = Pt(0)
            p_pr = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), SURFACE)
            p_pr.append(shading)
            add_inline(paragraph, " ".join(quote_parts))
            continue

        if stripped.startswith("*Nota.*"):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Figure Note")
            add_inline(paragraph, stripped)
            index += 1
            continue

        paragraph_buffer.append(raw)
        index += 1

    flush_paragraph()
    if in_code:
        add_code_block(document, code_lines)


def extract_segments(lines: list[str]) -> tuple[list[str], list[str]]:
    criterion_start = next(
        i
        for i, line in enumerate(lines)
        if line.startswith("## Criterio de evidencia")
        or line.startswith("## Alcance y criterios de validación")
    )
    index_start = next(i for i, line in enumerate(lines) if line.startswith("# Índice general"))
    main_start = next(i for i, line in enumerate(lines) if line.startswith("# 1. Introducción"))
    checklist_start = next(
        (i for i, line in enumerate(lines) if line.startswith("## Lista de control editorial")),
        len(lines),
    )
    criterion = lines[criterion_start:index_start]
    main = lines[main_start:checklist_start]
    return criterion, main


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    criterion, main = extract_segments(lines)

    document = Document()
    configure_styles(document)
    configure_document(document)
    add_cover(document)
    parse_markdown(document, criterion)
    document.add_page_break()
    add_indexes(document)
    parse_markdown(document, main)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
